"""
Divergências de conferência – back end (Flask).

Rotas:
  GET  /                  -> página (static/index.html)
  POST /api/ocr           -> recebe o print e devolve NF, fornecedor e itens divergentes
  GET  /api/memoria       -> empresas, vendedores e fornecedores já usados (autocompletar)
  POST /api/salvar        -> grava a divergência na planilha (.xlsx)
"""
import io
import json
import os
import sqlite3
import re
import shutil
import subprocess
import sys
import tempfile
import threading
import webbrowser
from datetime import date, datetime
from pathlib import Path

from flask import Flask, jsonify, request, send_from_directory
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from PIL import Image, ImageOps, ImageEnhance, ImageFilter

BASE = Path(__file__).resolve().parent
DATA = BASE / "data"
DATA.mkdir(exist_ok=True)

DB_PATH = DATA / "divergencias.db"


def inicializar_banco():
    with sqlite3.connect(DB_PATH) as conn:
        conn.execute("""
            CREATE TABLE IF NOT EXISTS divergencias (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                data_registro TEXT NOT NULL,
                empresa TEXT NOT NULL,
                vendedor TEXT DEFAULT '',
                nf TEXT DEFAULT '',
                fornecedor TEXT DEFAULT '',
                status TEXT NOT NULL DEFAULT 'PENDENTE',
                solucao TEXT DEFAULT '',
                dados_json TEXT NOT NULL
            )
        """)

        conn.execute("""
            CREATE INDEX IF NOT EXISTS idx_divergencias_empresa
            ON divergencias (empresa)
        """)

        conn.execute("""
            CREATE INDEX IF NOT EXISTS idx_divergencias_status
            ON divergencias (status)
        """)


inicializar_banco()

# Caminho da planilha. Mude aqui ou defina a variável de ambiente DIVERGENCIAS_XLSX.
XLSX_PATH = Path(
    os.environ.get(
        "DIVERGENCIAS_XLSX",
        r"C:\Users\estoque 3\Desktop\divergencias-app\divergenciasApp\data\divergencias 2026.xlsx"
    )
)
MEMORIA_PATH = DATA / "memoria.json"

MESES = ["JAN", "FEV", "MAR", "ABR", "MAI", "JUN", "JUL", "AGO", "SET", "OUT", "NOV", "DEZ"]
CABECALHO = ["EMPRESA / NF", "DATA", "VENDEDOR", "STATUS", "PAGAMENTO", "OBSERVAÇÃO"]
MOTIVOS = {
    "falta": "FALTA",
    "danificado": "DANIFICADO",
    "sem_validade": "SEM VALIDADE",
    "codigo_incorreto": "CÓDIGO INCORRETO",
    "sobra": "SOBRA",
}

RECURSOS = Path(__file__).resolve().parent / "static"

app = Flask(
    __name__,
    static_folder=str(RECURSOS),
    static_url_path="/static"
)
_lock = threading.Lock()


# --------------------------------------------------------------------------- OCR
def achar_tesseract():
    cmd = os.environ.get("TESSERACT_CMD") or shutil.which("tesseract")
    if cmd:
        return cmd
    for p in (
        r"C:\Program Files\Tesseract-OCR\tesseract.exe",
        r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
    ):
        if os.path.exists(p):
            return p
    return None


def rodar_ocr(img: Image.Image, psm: int) -> str:
    cmd = achar_tesseract()
    if not cmd:
        raise RuntimeError(
            "Tesseract OCR não encontrado. Instale-o (veja o README) "
            "ou defina a variável TESSERACT_CMD."
        )
    with tempfile.TemporaryDirectory() as tmp:
        arq = Path(tmp) / "print.png"
        img.save(arq)
        r = subprocess.run(
            [cmd, str(arq), "stdout", "--psm", str(psm)],
            capture_output=True, text=True, encoding="utf-8", errors="ignore",
        )
        return r.stdout


def preparar_simples(img: Image.Image) -> Image.Image:
    """Pré-processamento leve (o que já lia bem as sobras)."""
    img = ImageOps.exif_transpose(img).convert("L")
    if img.width < 2400:
        f = 2 if img.width < 1600 else 1.5
        img = img.resize((int(img.width * f), int(img.height * f)), Image.LANCZOS)
    return ImageOps.autocontrast(img)


def preparar_reforcado(img: Image.Image) -> Image.Image:
    """Pré-processamento com mais contraste/nitidez (ajuda em textos pequenos)."""
    img = ImageOps.exif_transpose(img).convert("L")

    if img.width < 2000:
        fator = 2000 / img.width
        img = img.resize(
            (2000, int(img.height * fator)),
            Image.Resampling.LANCZOS
        )

    img = ImageOps.autocontrast(img, cutoff=1)
    img = ImageEnhance.Contrast(img).enhance(1.5)
    img = ImageEnhance.Sharpness(img).enhance(1.5)
    img = img.filter(ImageFilter.UnsharpMask(radius=1, percent=130, threshold=3))
    return img


# Os números de quantidade aceitam O/o/I/l/S porque o OCR costuma confundir
# o "0" (ex.: Lançada: 0 numa sobra) com a letra O.
NUM = r"([\dOoIlS]{1,5})"

RE_ITEM = re.compile(
    r"Qtd\.?\s*[A-Za-z]{3,10}\s*[:;.]?\s*" + NUM + r"\s*[-–—]\s*[A-Za-z]{3,12}\s*[:;.]?\s*" + NUM + r"\s*"
    r"Prod\w*\s*[:;.]?\s*([\d ]{8,18}?)\s*[-–—]\s*(.+)",
    re.IGNORECASE,
)
RE_NF = re.compile(r"N[uúü]mero\W{0,6}(\d{3,10})", re.IGNORECASE)
RE_FORN = re.compile(r"F[oó0]\w{0,4}\W{1,3}(?:\S{1,2}\s+)?(\d{3,8})", re.IGNORECASE)

_TROCA_NUM = str.maketrans({"O": "0", "o": "0", "I": "1", "l": "1", "S": "5"})


def para_inteiro(txt: str) -> int:
    return int(txt.translate(_TROCA_NUM))


def corrigir_nome_produto(nome: str) -> str:
    """Corrige '4'->'A' e '0'->'O' apenas dentro de palavras (entre letras),
    sem mexer em medidas como 200ML, 45G, 100G."""
    def trocar(m):
        palavra = m.group(0)
        # palavra com número de 2+ dígitos é medida/código: preserva
        if re.search(r"\d{2,}", palavra):
            return palavra
        palavra = re.sub(r"(?<=[A-Za-zÀ-ÿ])4(?=[A-Za-zÀ-ÿ])", "A", palavra)
        palavra = re.sub(r"(?<=[A-Za-zÀ-ÿ])0(?=[A-Za-zÀ-ÿ])", "O", palavra)
        return palavra

    return re.sub(r"\S+", trocar, nome.strip())


def interpretar(texto: str):
    itens = []
    for linha in texto.splitlines():
        m = RE_ITEM.search(linha)
        if not m:
            continue
        lanc, conf, cod, nome = m.groups()
        cod = re.sub(r"\D", "", cod)
        nome = re.sub(r"[\s,.;|]+$", "", nome).strip()
        nome = corrigir_nome_produto(nome)
        try:
            lanc_i = para_inteiro(lanc)
            conf_i = para_inteiro(conf)
        except ValueError:
            continue
        itens.append({
            "codigo": cod,
            "nome": nome,
            "lancada": lanc_i,
            "conferida": conf_i,
        })
    nf = RE_NF.search(texto)
    forn = RE_FORN.search(texto)
    return itens, (nf.group(1) if nf else ""), (forn.group(1) if forn else "")


def ean13_valido(cod: str) -> bool:
    if len(cod) != 13 or not cod.isdigit():
        return False
    soma = sum(int(c) * (3 if i % 2 else 1) for i, c in enumerate(cod[:12]))
    return (10 - soma % 10) % 10 == int(cod[12])


def _distancia(a: str, b: str) -> int:
    if len(a) != len(b):
        return 99
    return sum(x != y for x, y in zip(a, b))


def agrupar_itens(candidatos):
    """Une leituras do mesmo item feitas em passagens diferentes do OCR.
    Mesmo item = mesmas quantidades e código quase igual (até 2 dígitos
    trocados). Dentro do grupo, escolhe o código com dígito verificador
    EAN-13 válido; se não houver, o mais repetido."""
    grupos = []
    for it in candidatos:
        for g in grupos:
            base = g[0]
            if (base["lancada"] == it["lancada"]
                    and base["conferida"] == it["conferida"]
                    and _distancia(base["codigo"], it["codigo"]) <= 2):
                g.append(it)
                break
        else:
            grupos.append([it])

    resultado = []
    for g in grupos:
        contagem = {}
        for it in g:
            contagem[it["codigo"]] = contagem.get(it["codigo"], 0) + 1
        codigos = sorted(contagem, key=lambda c: (ean13_valido(c), contagem[c]), reverse=True)
        cod = codigos[0]
        do_codigo = [it for it in g if it["codigo"] == cod]
        nomes = {}
        for it in do_codigo:
            nomes[it["nome"]] = nomes.get(it["nome"], 0) + 1
        nome = max(nomes, key=nomes.get)
        resultado.append({
            "codigo": cod,
            "nome": nome,
            "lancada": g[0]["lancada"],
            "conferida": g[0]["conferida"],
        })
    return resultado


@app.post("/api/ocr")
def api_ocr():
    arq = request.files.get("imagem")

    if not arq:
        return jsonify(erro="Nenhuma imagem enviada."), 400

    try:
        original = Image.open(io.BytesIO(arq.read()))
        original.load()

        # Roda as duas preparações × os dois modos de leitura e JUNTA os itens
        # encontrados em cada tentativa (por código). Assim, um item (como uma
        # sobra) que uma leitura perdeu é recuperado pela outra.
        candidatos = []
        nf_final = ""
        forn_final = ""
        texto_final = ""

        for preparo in (preparar_simples, preparar_reforcado):
            img = preparo(original)
            for psm in (11, 6):
                bruto = rodar_ocr(img, psm)
                itens_lidos, nf_lida, forn_lido = interpretar(bruto)

                if len(bruto) > len(texto_final):
                    texto_final = bruto
                nf_final = nf_final or nf_lida
                forn_final = forn_final or forn_lido
                candidatos.extend(itens_lidos)

        melhor_itens = agrupar_itens(candidatos)

    except RuntimeError as e:
        return jsonify(erro=str(e)), 500
    except Exception as e:
        return jsonify(erro=f"Não consegui ler a imagem: {e}"), 400

    mem = ler_memoria()

    return jsonify(
        nf=nf_final,
        fornecedor=forn_final,
        empresa_sugerida=mem["fornecedores"].get(forn_final, ""),
        itens=melhor_itens,
        texto_bruto=texto_final,
    )


# ------------------------------------------------------------------- memória
def ler_memoria():
    base = {"empresas": [], "vendedores": [], "fornecedores": {}}
    if MEMORIA_PATH.exists():
        try:
            base.update(json.loads(MEMORIA_PATH.read_text(encoding="utf-8")))
        except Exception:
            pass
    return base



def gravar_memoria(empresa, vendedor, fornecedor):
    mem = ler_memoria()

    # Garante que as estruturas existam
    mem.setdefault("empresas", [])
    mem.setdefault("vendedores", [])
    mem.setdefault("fornecedores", {})

    empresa = str(empresa or "").strip().upper()
    vendedor = str(vendedor or "").strip().upper()
    fornecedor = str(fornecedor or "").strip()

    if empresa and empresa not in mem["empresas"]:
        mem["empresas"].append(empresa)

    if vendedor and vendedor not in mem["vendedores"]:
        mem["vendedores"].append(vendedor)

    if fornecedor and empresa:
        mem["fornecedores"][fornecedor] = empresa

    mem["empresas"].sort()
    mem["vendedores"].sort()

    # Gravação segura do JSON
    temporario = MEMORIA_PATH.with_suffix(".tmp")

    temporario.write_text(
        json.dumps(mem, ensure_ascii=False, indent=2),
        encoding="utf-8"
    )

    temporario.replace(MEMORIA_PATH)



@app.get("/api/memoria")
def api_memoria():
    m = ler_memoria()
    m["caminho_planilha"] = str(XLSX_PATH)
    return jsonify(m)


# ------------------------------------------------------------------- planilha
BORDA = Border(*(Side(style="thin"),) * 4)
CINZA = PatternFill("solid", fgColor="999999")
CINZA_CAB = PatternFill("solid", fgColor="CCCCCC")
VERDE = PatternFill("solid", fgColor="99FF66")
BRANCO = PatternFill("solid", fgColor="FFFFFF")


def criar_planilha(caminho: Path):
    wb = Workbook()
    wb.remove(wb.active)
    larguras = [46, 14, 24, 18, 26, 26]
    for mes in MESES:
        ws = wb.create_sheet(mes)
        ws.merge_cells("A1:F1")
        ws["A1"].fill = CINZA
        for col, titulo in enumerate(CABECALHO, 1):
            c = ws.cell(row=2, column=col, value=titulo)
            c.font = Font(name="Arial", bold=True, size=10)
            c.fill = CINZA_CAB
            c.border = BORDA
            c.alignment = Alignment(horizontal="center", vertical="center")
            ws.column_dimensions[c.column_letter].width = larguras[col - 1]
        dv = DataValidation(type="list", formula1='"OK,PENDENTE"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add("D3:D1000")
        ws.freeze_panes = "A3"
    wb.save(caminho)


def estilo_padrao(cel):
    cel.font = Font(name="Arial", size=10)
    cel.border = BORDA
    cel.alignment = Alignment(horizontal="center", vertical="center")


def primeira_linha_livre(ws):
    linha = 3
    while ws.cell(row=linha, column=1).value not in (None, ""):
        linha += 1
    return linha


def gravar_planilha(d):
    dt = datetime.strptime(d["data"], "%Y-%m-%d").date() if d.get("data") else date.today()
    empresa = d["empresa"].strip().upper()
    nf = str(d["nf"]).strip()
    vendedor = d.get("vendedor", "").strip().upper()
    nome_aba = MESES[dt.month - 1]

    with _lock:
        if not XLSX_PATH.exists():
            XLSX_PATH.parent.mkdir(parents=True, exist_ok=True)
            criar_planilha(XLSX_PATH)
        wb = load_workbook(XLSX_PATH)
        if nome_aba not in wb.sheetnames:
            wb.create_sheet(nome_aba)
        ws = wb[nome_aba]

        # evita lançar a mesma NF duas vezes
        if not d.get("forcar"):
            for (v,) in ws.iter_rows(min_row=3, max_col=1, values_only=True):
                if v and str(v).strip().endswith(nf):
                    return {"duplicada": True, "existente": str(v)}

        linha = primeira_linha_livre(ws)
        valores = [f"{empresa} {nf}", dt, vendedor, "PENDENTE", None, None]
        for col, val in enumerate(valores, 1):
            cel = ws.cell(row=linha, column=col)
            if linha > 3:
                cel._style = ws.cell(row=linha - 1, column=col)._style
            else:
                estilo_padrao(cel)
            cel.value = val
        ws.cell(row=linha, column=2).number_format = "dd/mm/yy"
        ws.cell(row=linha, column=4).fill = BRANCO
        # a linha anterior pode estar verde ("OK"); aqui sempre nasce pendente.
        ws.cell(row=linha, column=5).fill = PatternFill(fill_type=None)
        ws.cell(row=linha, column=6).fill = PatternFill(fill_type=None)

        # detalhes dos itens num comentário da célula A (o triângulo rosa do Calc)
        linhas = []
        for it in d["itens"]:
            motivo = MOTIVOS.get(it["motivo"], it["motivo"])
            extra = ""
            if it["motivo"] == "codigo_incorreto":
                extra = f" (faturado {it.get('cod_errado','')} -> correto {it.get('cod_correto','')})"
            linhas.append(f"{it['nome']} [{it['codigo']}] – {motivo} {it['qtd']}un{extra}")
        com = Comment("\n".join(linhas), "Divergências")
        com.width, com.height = 420, 40 + 18 * len(linhas)
        ws.cell(row=linha, column=1).comment = com

        wb.save(XLSX_PATH)
    return {"duplicada": False, "aba": nome_aba, "linha": linha}


@app.post("/api/salvar")
def api_salvar():
    d = request.get_json(force=True)

    if not d.get("empresa") or not d.get("nf") or not d.get("itens"):
        return jsonify(erro="Faltam dados (empresa, NF ou itens)."), 400

    try:
        r = gravar_planilha(d)
    except PermissionError:
        return jsonify(
            erro="Não consegui gravar: feche a planilha no LibreOffice/Excel e tente de novo."
        ), 423
    except Exception as e:
        return jsonify(erro=f"Erro ao gravar a planilha: {e}"), 500

    if not r["duplicada"]:
        empresa = d["empresa"].strip().upper()
        vendedor = str(d.get("vendedor", "")).strip().upper()
        fornecedor = str(d.get("fornecedor", "")).strip()

        try:
            data_registro = (
                d.get("data")
                or d.get("data_conferencia")
                or datetime.now().strftime("%Y-%m-%d")
            )

            with sqlite3.connect(DB_PATH) as conn:
                cursor = conn.execute(
                    """
                    INSERT INTO divergencias (
                        data_registro,
                        empresa,
                        vendedor,
                        nf,
                        fornecedor,
                        status,
                        solucao,
                        dados_json
                    )
                    VALUES (?, ?, ?, ?, ?, 'PENDENTE', '', ?)
                    """,
                    (
                        data_registro,
                        empresa,
                        vendedor,
                        str(d.get("nf", "")),
                        fornecedor,
                        json.dumps(d, ensure_ascii=False),
                    ),
                )
                id_divergencia = cursor.lastrowid

                       # Atualiza a memória após gravar no banco
            gravar_memoria(
                empresa,
                vendedor,
                fornecedor
            )

            app.logger.info(
                "Divergência %s gravada no DB e memória atualizada.",
                id_divergencia
            )
            
        except Exception as e:
            app.logger.exception("Erro ao salvar divergência no SQLite")
            return jsonify(
                erro=(
                    "A planilha foi gravada, mas houve um erro ao salvar "
                    f"no banco de dados: {e}"
                ),
                planilha_gravada=True,
            ), 500

        r["id_divergencia"] = id_divergencia

    r["caminho"] = str(XLSX_PATH)
    return jsonify(r)


@app.get("/api/divergencias")
def api_listar_divergencias():
    data_inicio = request.args.get("data_inicio", "").strip()
    data_fim = request.args.get("data_fim", "").strip()
    empresa = request.args.get("empresa", "").strip()
    vendedor = request.args.get("vendedor", "").strip()
    status = request.args.get("status", "").strip().upper()

    sql = "SELECT * FROM divergencias WHERE 1=1"
    parametros = []

    if data_inicio:
        sql += " AND date(data_registro) >= date(?)"
        parametros.append(data_inicio)

    if data_fim:
        sql += " AND date(data_registro) <= date(?)"
        parametros.append(data_fim)

    if empresa:
        sql += " AND empresa LIKE ?"
        parametros.append(f"%{empresa}%")

    if vendedor:
        sql += " AND vendedor LIKE ?"
        parametros.append(f"%{vendedor}%")

    if status in ("PENDENTE", "OK"):
        sql += " AND status = ?"
        parametros.append(status)

    sql += " ORDER BY id DESC"

    try:
        with sqlite3.connect(DB_PATH) as conn:
            conn.row_factory = sqlite3.Row
            registros = conn.execute(sql, parametros).fetchall()

        resultado = []
        for registro in registros:
            item = dict(registro)
            try:
                item["dados"] = json.loads(item.pop("dados_json"))
            except (ValueError, TypeError):
                item["dados"] = {}

            resultado.append(item)

        return jsonify(resultado)

    except Exception as e:
        app.logger.exception("Erro ao consultar divergências")
        return jsonify(erro=f"Erro ao consultar divergências: {e}"), 500


@app.delete("/api/divergencias/<int:id_divergencia>")
def api_excluir_divergencia(id_divergencia):
    try:
        with sqlite3.connect(DB_PATH) as conn:
            cursor = conn.execute(
                """
                DELETE FROM divergencias
                WHERE id = ?
                """,
                (id_divergencia,)
            )

            if cursor.rowcount == 0:
                return jsonify(
                    erro="Divergência não encontrada."
                ), 404

        return jsonify(
            sucesso=True,
            mensagem="Divergência excluída com sucesso.",
            id=id_divergencia
        )

    except Exception as e:
        app.logger.exception("Erro ao excluir divergência")

        return jsonify(
            erro=f"Erro ao excluir divergência: {e}"
        ), 500


@app.post("/api/divergencias/<int:id_divergencia>/status")
def api_atualizar_status_divergencia(id_divergencia):
    d = request.get_json(silent=True) or {}

    status = str(d.get("status", "")).strip().upper()
    solucao = str(d.get("solucao", "")).strip()

    if status not in ("PENDENTE", "OK"):
        return jsonify(erro="Status inválido."), 400

    if status == "OK" and not solucao:
        return jsonify(
            erro="Descreva como a divergência foi resolvida."
        ), 400

    try:
        with sqlite3.connect(DB_PATH) as conn:
            cursor = conn.execute(
                """
                UPDATE divergencias
                SET status = ?, solucao = ?
                WHERE id = ?
                """,
                (status, solucao, id_divergencia),
            )

            if cursor.rowcount == 0:
                return jsonify(erro="Divergência não encontrada."), 404

        return jsonify(
            sucesso=True,
            id=id_divergencia,
            status=status,
            solucao=solucao,
        )

    except Exception as e:
        app.logger.exception("Erro ao atualizar divergência")
        return jsonify(erro=f"Erro ao atualizar status: {e}"), 500


# --------------------------------------------------------------------- páginas
@app.get("/")
def index():
    return send_from_directory(app.static_folder, "index.html")


if __name__ == "__main__":
    porta = int(os.environ.get("PORT", 5000))
    if "--no-browser" not in sys.argv:
        threading.Timer(1.2, lambda: webbrowser.open(f"http://127.0.0.1:{porta}")).start()
    app.run(host="127.0.0.1", port=porta, debug=False)