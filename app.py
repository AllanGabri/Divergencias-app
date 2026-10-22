"""
Divergências de conferência – back end (Flask).

Rotas:
  GET  /                  -> página (static/index.html)
  POST /api/ocr           -> recebe o print e devolve NF, fornecedor e itens divergentes
  GET  /api/memoria       -> empresas, vendedores e fornecedores já usados (autocompletar)
  POST /api/salvar        -> grava a divergência na planilha (.xlsx)
"""
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import threading
import webbrowser
from datetime import date, datetime
from pathlib import Path

from flask import Flask, jsonify, request, send_from_directory
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from PIL import Image, ImageOps

BASE = Path(__file__).resolve().parent
DATA = BASE / "data"
DATA.mkdir(exist_ok=True)

# Caminho da planilha. Mude aqui ou defina a variável de ambiente DIVERGENCIAS_XLSX.
XLSX_PATH = Path(
    os.environ.get(
        "DIVERGENCIAS_XLSX",
        r"C:\Users\estoque 3\Desktop\divergencias-app\divergenciasApp\data\divergencias 2026.xlsx"
    )
)
MEMORIA_PATH = DATA / "memoria.json"

MESES = ["JAN", "FEV", "MAR", "ABR", "MAI", "JUN", "JUL", "AGO", "SET", "OUT", "NOV", "DEZ"]
CABECALHO = ["EMPRESA / NF", "DATA", "VENDEDOR", "STATUS", "PAGAMENTO", "OBSERVAÇÃO"]
MOTIVOS = {
    "falta": "FALTA",
    "danificado": "DANIFICADO",
    "sem_validade": "SEM VALIDADE",
    "codigo_incorreto": "CÓDIGO INCORRETO",
    "sobra": "SOBRA",
}

RECURSOS = Path(__file__).resolve().parent / "static"

app = Flask(
    __name__,
    static_folder=str(RECURSOS),
    static_url_path="/static"
)
_lock = threading.Lock()


# --------------------------------------------------------------------------- OCR
def achar_tesseract():
    cmd = os.environ.get("TESSERACT_CMD") or shutil.which("tesseract")
    if cmd:
        return cmd
    for p in (
        r"C:\Program Files\Tesseract-OCR\tesseract.exe",
        r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
    ):
        if os.path.exists(p):
            return p
    return None


def rodar_ocr(img: Image.Image, psm: int) -> str:
    cmd = achar_tesseract()
    if not cmd:
        raise RuntimeError(
            "Tesseract OCR não encontrado. Instale-o (veja o README) "
            "ou defina a variável TESSERACT_CMD."
        )
    with tempfile.TemporaryDirectory() as tmp:
        arq = Path(tmp) / "print.png"
        img.save(arq)
        r = subprocess.run(
            [cmd, str(arq), "stdout", "--psm", str(psm)],
            capture_output=True, text=True, encoding="utf-8", errors="ignore",
        )
        return r.stdout


def preparar(img: Image.Image) -> Image.Image:
    img = ImageOps.exif_transpose(img).convert("L")
    if img.width < 2400:
        f = 2 if img.width < 1600 else 1.5
        img = img.resize((int(img.width * f), int(img.height * f)), Image.LANCZOS)
    return ImageOps.autocontrast(img)


RE_ITEM = re.compile(
    r"Qtd\.?\s*[A-Za-z]{3,10}\s*:?\s*(\d+)\s*[-–—]\s*[A-Za-z]{3,12}\s*:?\s*(\d+)\s*"
    r"Prod\w*\s*:?\s*([\d ]{8,18}?)\s*[-–—]\s*(.+)",
    re.IGNORECASE,
)
RE_NF = re.compile(r"N[uúü]mero\W{0,6}(\d{3,10})", re.IGNORECASE)
RE_FORN = re.compile(r"F[oó0]\w{0,4}\W{1,3}(?:\S{1,2}\s+)?(\d{3,8})", re.IGNORECASE)


def interpretar(texto: str):
    itens = []
    for linha in texto.splitlines():
        m = RE_ITEM.search(linha)
        if not m:
            continue
        lanc, conf, cod, nome = m.groups()
        cod = re.sub(r"\D", "", cod)
        nome = re.sub(r"[\s,.;|]+$", "", nome).strip()
        itens.append({
            "codigo": cod,
            "nome": nome,
            "lancada": int(lanc),
            "conferida": int(conf),
        })
    nf = RE_NF.search(texto)
    forn = RE_FORN.search(texto)
    return itens, (nf.group(1) if nf else ""), (forn.group(1) if forn else "")


@app.post("/api/ocr")
def api_ocr():
    arq = request.files.get("imagem")
    if not arq:
        return jsonify(erro="Nenhuma imagem enviada."), 400
    try:
        img = preparar(Image.open(io.BytesIO(arq.read())))
        itens, nf, forn, bruto = [], "", "", ""
        # psm 11 (texto esparso) costuma ler melhor essa tela; psm 6 fica de reserva
        for psm in (11, 6):
            bruto = rodar_ocr(img, psm)
            i, n, f = interpretar(bruto)
            if len(i) > len(itens):
                itens = i
            nf = nf or n
            forn = forn or f
            if itens and nf:
                break
    except RuntimeError as e:
        return jsonify(erro=str(e)), 500
    except Exception as e:  # imagem inválida etc.
        return jsonify(erro=f"Não consegui ler a imagem: {e}"), 400

    mem = ler_memoria()
    return jsonify(
        nf=nf,
        fornecedor=forn,
        empresa_sugerida=mem["fornecedores"].get(forn, ""),
        itens=itens,
        texto_bruto=bruto,
    )


# ------------------------------------------------------------------- memória
def ler_memoria():
    base = {"empresas": [], "vendedores": [], "fornecedores": {}}
    if MEMORIA_PATH.exists():
        try:
            base.update(json.loads(MEMORIA_PATH.read_text(encoding="utf-8")))
        except Exception:
            pass
    return base


def gravar_memoria(empresa, vendedor, fornecedor):
    mem = ler_memoria()
    if empresa and empresa not in mem["empresas"]:
        mem["empresas"].append(empresa)
    if vendedor and vendedor not in mem["vendedores"]:
        mem["vendedores"].append(vendedor)
    if fornecedor and empresa:
        mem["fornecedores"][fornecedor] = empresa
    mem["empresas"].sort()
    mem["vendedores"].sort()
    MEMORIA_PATH.write_text(json.dumps(mem, ensure_ascii=False, indent=2), encoding="utf-8")


@app.get("/api/memoria")
def api_memoria():
    m = ler_memoria()
    m["caminho_planilha"] = str(XLSX_PATH)
    return jsonify(m)


# ------------------------------------------------------------------- planilha
BORDA = Border(*(Side(style="thin"),) * 4)
CINZA = PatternFill("solid", fgColor="999999")
CINZA_CAB = PatternFill("solid", fgColor="CCCCCC")
VERDE = PatternFill("solid", fgColor="99FF66")
AMARELO = PatternFill("solid", fgColor="FFFF66")


def criar_planilha(caminho: Path):
    wb = Workbook()
    wb.remove(wb.active)
    larguras = [46, 14, 24, 18, 26, 26]
    for mes in MESES:
        ws = wb.create_sheet(mes)
        ws.merge_cells("A1:F1")
        ws["A1"].fill = CINZA
        for col, titulo in enumerate(CABECALHO, 1):
            c = ws.cell(row=2, column=col, value=titulo)
            c.font = Font(name="Arial", bold=True, size=10)
            c.fill = CINZA_CAB
            c.border = BORDA
            c.alignment = Alignment(horizontal="center", vertical="center")
            ws.column_dimensions[c.column_letter].width = larguras[col - 1]
        dv = DataValidation(type="list", formula1='"OK,PENDENTE"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add("D3:D1000")
        ws.freeze_panes = "A3"
    wb.save(caminho)


def estilo_padrao(cel):
    cel.font = Font(name="Arial", size=10)
    cel.border = BORDA
    cel.alignment = Alignment(horizontal="center", vertical="center")


def primeira_linha_livre(ws):
    linha = 3
    while ws.cell(row=linha, column=1).value not in (None, ""):
        linha += 1
    return linha


def gravar_planilha(d):
    dt = datetime.strptime(d["data"], "%Y-%m-%d").date() if d.get("data") else date.today()
    empresa = d["empresa"].strip().upper()
    nf = str(d["nf"]).strip()
    vendedor = d.get("vendedor", "").strip().upper()
    nome_aba = MESES[dt.month - 1]

    with _lock:
        if not XLSX_PATH.exists():
            XLSX_PATH.parent.mkdir(parents=True, exist_ok=True)
            criar_planilha(XLSX_PATH)
        wb = load_workbook(XLSX_PATH)
        if nome_aba not in wb.sheetnames:
            wb.create_sheet(nome_aba)
        ws = wb[nome_aba]

        # evita lançar a mesma NF duas vezes
        if not d.get("forcar"):
            for (v,) in ws.iter_rows(min_row=3, max_col=1, values_only=True):
                if v and str(v).strip().endswith(nf):
                    return {"duplicada": True, "existente": str(v)}

        linha = primeira_linha_livre(ws)
        valores = [f"{empresa} {nf}", dt, vendedor, "PENDENTE", None, None]
        for col, val in enumerate(valores, 1):
            cel = ws.cell(row=linha, column=col)
            if linha > 3:
                cel._style = ws.cell(row=linha - 1, column=col)._style
            else:
                estilo_padrao(cel)
            cel.value = val
        ws.cell(row=linha, column=2).number_format = "dd/mm/yy"
        ws.cell(row=linha, column=4).fill = AMARELO
        # a linha anterior pode estar verde ("OK"); aqui sempre nasce pendente.
        ws.cell(row=linha, column=5).fill = PatternFill(fill_type=None)
        ws.cell(row=linha, column=6).fill = PatternFill(fill_type=None)

        # detalhes dos itens num comentário da célula A (o triângulo rosa do Calc)
        linhas = []
        for it in d["itens"]:
            motivo = MOTIVOS.get(it["motivo"], it["motivo"])
            extra = ""
            if it["motivo"] == "codigo_incorreto":
                extra = f" (faturado {it.get('cod_errado','')} -> correto {it.get('cod_correto','')})"
            linhas.append(f"{it['nome']} [{it['codigo']}] – {motivo} {it['qtd']}un{extra}")
        com = Comment("\n".join(linhas), "Divergências")
        com.width, com.height = 420, 40 + 18 * len(linhas)
        ws.cell(row=linha, column=1).comment = com

        wb.save(XLSX_PATH)
    return {"duplicada": False, "aba": nome_aba, "linha": linha}


@app.post("/api/salvar")
def api_salvar():
    d = request.get_json(force=True)
    if not d.get("empresa") or not d.get("nf") or not d.get("itens"):
        return jsonify(erro="Faltam dados (empresa, NF ou itens)."), 400
    try:
        r = gravar_planilha(d)
    except PermissionError:
        return jsonify(erro="Não consegui gravar: feche a planilha no LibreOffice/Excel e tente de novo."), 423
    except Exception as e:
        return jsonify(erro=f"Erro ao gravar a planilha: {e}"), 500
    if not r["duplicada"]:
        gravar_memoria(d["empresa"].strip().upper(), d.get("vendedor", "").strip().upper(), str(d.get("fornecedor", "")).strip())
    r["caminho"] = str(XLSX_PATH)
    return jsonify(r)


# --------------------------------------------------------------------- páginas
@app.get("/")
def index():
    return send_from_directory(app.static_folder, "index.html")


if __name__ == "__main__":
    porta = int(os.environ.get("PORT", 5000))
    if "--no-browser" not in sys.argv:
        threading.Timer(1.2, lambda: webbrowser.open(f"http://127.0.0.1:{porta}")).start()
    app.run(host="127.0.0.1", port=porta, debug=False)
